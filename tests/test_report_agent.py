"""
Tests for the Report Agent (Phase 5).

No test needs a Gemini API key or network access. Reports are written to
temporary directories. Run with:  python -m pytest -q
"""

import copy
import datetime
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document

from agents import report_agent as ra
from agents.report_agent import (
    build_references,
    format_report_markdown,
    generate_research_report,
    save_report_docx,
    build_report_blocks,
)

PAPER_A = "https://openalex.org/W1"
PAPER_B = "https://openalex.org/W2"
PAPER_C = "https://doi.org/10.1000/xyz"

QUESTION = "How does federated learning protect privacy?"
TODAY = datetime.date(2026, 10, 20)


def make_papers():
    return [
        {"id": PAPER_A, "title": "Federated Averaging", "year": 2017,
         "authors": ["Ann One", "Bob Two"], "doi": "https://doi.org/10.1234/a",
         "url": "https://doi.org/10.1234/a",
         "abstract": "Federated learning keeps data on device.",
         "source_questions": ["What is federated learning?"]},
        {"id": PAPER_B, "title": "Secure Aggregation", "year": 2019,
         "authors": ["Cy Three", "Di Four", "Ed Five", "Fi Six"],
         "doi": None, "url": PAPER_B,
         "abstract": "Secure aggregation hides individual updates.",
         "source_questions": ["What is federated learning?",
                              "How is privacy protected?"]},
        {"id": PAPER_C, "title": "Crossref Only Record", "year": None,
         "authors": [], "doi": "https://doi.org/10.1000/xyz",
         "url": PAPER_C, "abstract": "",
         "source_questions": ["How is privacy protected?"]},
    ]


def make_plan():
    return {
        "main_topic": "Federated learning privacy",
        "research_objective": "Understand privacy protection in FL.",
        "research_plan": ["Define FL", "Survey privacy methods"],
        "sub_questions": ["What is federated learning?",
                          "How is privacy protected?"],
    }


def make_graph():
    return [{
        "title": "Federated Averaging", "year": 2017,
        "authors": ["Ann One"], "url": "https://doi.org/10.1234/a",
        "relevance_score": 3.0, "graph_score": 2.0,
        "combined_graph_score": 4.0,
        "matched_concepts": ["federated", "privacy"],
        "graph_sources": [],
        "related_papers": [{"title": "Secure Aggregation", "year": 2019,
                            "url": PAPER_B, "shared_concepts": 3}],
    }]


def verification(status, confidence=0.8, excerpt="", warnings=None,
                 needs_review=None):
    return {
        "status": status, "confidence": confidence,
        "justification": f"Judged {status}.",
        "evidence_excerpt": excerpt, "evidence_paper_id": PAPER_A,
        "evidence_paper_title": "Federated Averaging",
        "warnings": warnings or [],
        "needs_review": (status != "SUPPORTED") if needs_review is None
        else needs_review,
    }


def make_claim(cid, status, **kwargs):
    claim = {
        "claim_id": cid, "claim": f"Claim text for {cid}.",
        "support_type": "source_supported", "paper_ids": [PAPER_A],
        "invalid_paper_ids": [], "evidence_excerpt": "keeps data on device",
        "excerpt_verified": True, "evidence_status": "excerpt_matched",
        "verification_note": "original",
    }
    if status:
        claim["verification"] = verification(status, **kwargs)
    return claim


def make_analysis(claims=None, verified=True):
    claims = claims if claims is not None else [
        make_claim("CLM-001", "SUPPORTED", excerpt="keeps data on device")]
    analysis = {
        "status": "success",
        "summary": "FL keeps data local and uses secure aggregation.",
        "key_themes_and_consensus": [{
            "theme": "Data locality",
            "description": "Data stays on the device.",
            "supporting_paper_ids": [PAPER_A], "invalid_paper_ids": []}],
        "divergent_findings_and_contradictions": [],
        "methodologies_and_benchmarks": [{
            "approach_or_dataset": "Secure aggregation",
            "description": "Hides updates.", "papers_using": [PAPER_B],
            "invalid_paper_ids": []}],
        "limitations_and_research_gaps": [
            {"gap_or_limitation": "Communication cost is high.",
             "reported_by": [PAPER_B], "type": "reported_limitation",
             "invalid_paper_ids": []},
            {"gap_or_limitation": "Long-term privacy is unstudied.",
             "reported_by": [], "type": "inferred_gap",
             "invalid_paper_ids": []}],
        "evidence_linked_claims": claims,
        "metadata": {},
    }
    if verified:
        analysis["verification_summary"] = {
            "service_status": "completed", "service_note": "",
            "total_claims": len(claims)}
    return analysis


def md(analysis=None, **kwargs):
    kwargs.setdefault("papers", make_papers())
    kwargs.setdefault("graph_context", make_graph())
    kwargs.setdefault("plan", make_plan())
    return format_report_markdown(
        QUESTION, analysis=analysis if analysis is not None
        else make_analysis(), generated_at=TODAY, **kwargs)


class TestMarkdown(unittest.TestCase):

    def test_1_sections_in_order(self):
        text = md()
        headings = [
            "1. Title and Research Metadata", "2. Executive Summary",
            "3. Research Plan and Sub-Questions",
            "4. Literature Landscape and Paper Graph",
            "5. Comparative Findings and Synthesis",
            "6. Contradictions and Open Research Gaps",
            "7. Verified Claims and Evidence Matrix",
            "8. Verification Summary and Limitations",
            "9. Conclusion and Future Research",
            "10. References and Source Index"]
        positions = [text.index(h) for h in headings]
        self.assertEqual(positions, sorted(positions))
        self.assertTrue(text.startswith(
            "# Research Report: How does federated learning protect "
            "privacy"))

    def test_metadata_counts(self):
        text = md()
        self.assertIn("2026-10-20", text)
        self.assertIn("| Retrieved papers \\(unique\\)" if False else
                      "Retrieved papers", text)
        blocks = build_report_blocks(
            QUESTION, make_plan(), make_papers(), make_graph(),
            make_analysis(), generated_at=TODAY)
        table = next(b for b in blocks if b["type"] == "table")
        values = {row[0]: row[1] for row in table["rows"]}
        self.assertEqual(values["Retrieved papers (unique)"], "3")
        self.assertEqual(values["Papers with abstracts"], "2")
        self.assertEqual(values["Analyzed claims"], "1")

    def test_plan_and_subquestions_listed(self):
        text = md()
        self.assertIn("What is federated learning?", text)
        self.assertIn("How is privacy protected?", text)
        self.assertIn("Survey privacy methods", text)

    def test_empty_question_rejected(self):
        with self.assertRaises(ValueError):
            format_report_markdown("  ", {}, [], [], {})

    def test_never_calls_whole_report_verified(self):
        text = md()
        self.assertIn("preliminary", text.lower())
        self.assertIn("NOT been independently verified", text)
        self.assertNotIn("fully verified", text.lower())


class TestDocx(unittest.TestCase):

    def test_2_docx_is_created_and_has_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = generate_research_report(
                QUESTION, make_plan(), make_papers(), make_graph(),
                make_analysis(), output_dir=tmp, generated_at=TODAY)
            self.assertEqual(result["errors"], [])
            self.assertTrue(Path(result["markdown_path"]).exists())
            self.assertTrue(Path(result["docx_path"]).exists())

            doc = Document(result["docx_path"])
            headings = [p.text for p in doc.paragraphs
                        if p.style.name.startswith("Heading 1")]
            self.assertEqual(len(headings), 10)
            self.assertEqual(headings[0], "1. Title and Research Metadata")
            self.assertGreaterEqual(len(doc.tables), 5)

    def test_docx_has_page_numbers_and_landscape_matrix(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = generate_research_report(
                QUESTION, make_plan(), make_papers(), make_graph(),
                make_analysis(), output_dir=tmp, generated_at=TODAY)
            doc = Document(result["docx_path"])
            self.assertTrue(any(
                s.page_width > s.page_height for s in doc.sections))
            with zipfile.ZipFile(result["docx_path"]) as archive:
                footers = "".join(
                    archive.read(n).decode("utf-8")
                    for n in archive.namelist() if "footer" in n)
            self.assertIn("PAGE", footers)

    def test_docx_urls_are_clickable(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = generate_research_report(
                QUESTION, make_plan(), make_papers(), make_graph(),
                make_analysis(), output_dir=tmp, generated_at=TODAY)
            with zipfile.ZipFile(result["docx_path"]) as archive:
                body = archive.read("word/document.xml").decode("utf-8")
                rels = archive.read(
                    "word/_rels/document.xml.rels").decode("utf-8")
            self.assertIn("<w:hyperlink", body)
            self.assertIn("https://openalex.org/W2", rels)

    def test_markdown_and_docx_have_the_same_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = generate_research_report(
                QUESTION, make_plan(), make_papers(), make_graph(),
                make_analysis(), output_dir=tmp, generated_at=TODAY)
            doc = Document(result["docx_path"])
            doc_text = " ".join(
                [p.text for p in doc.paragraphs]
                + [c.text for t in doc.tables for r in t.rows
                   for c in r.cells])
            markdown = Path(result["markdown_path"]).read_text(
                encoding="utf-8")
            for phrase in ("Data locality", "Secure aggregation",
                           "Communication cost is high.",
                           "Claim text for CLM-001."):
                self.assertIn(phrase, doc_text)
                self.assertIn(phrase, markdown)

    def test_output_directory_is_created(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "nested" / "output"
            result = generate_research_report(
                QUESTION, make_plan(), make_papers(), [],
                make_analysis(), output_dir=target, generated_at=TODAY)
            self.assertTrue((target / "report.md").exists())
            self.assertTrue((target / "report.docx").exists())
            self.assertEqual(result["errors"], [])

    def test_docx_failure_is_reported_not_hidden(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(ra, "save_report_docx",
                              side_effect=OSError("disk full")):
                result = generate_research_report(
                    QUESTION, make_plan(), make_papers(), [],
                    make_analysis(), output_dir=tmp)
            self.assertIsNone(result["docx_path"])
            self.assertTrue(result["markdown_path"])
            self.assertIn("Word report was not saved",
                          result["errors"][0])
            self.assertFalse((Path(tmp) / "report.docx").exists())


class TestGraph(unittest.TestCase):

    def test_3_graph_present_is_described(self):
        text = md()
        self.assertIn("not citation links", text)
        self.assertIn("federated, privacy", text)
        self.assertIn("3 shared concepts", text)

    def test_3_graph_empty_or_missing_is_stated(self):
        for graph in (None, [], "garbage"):
            text = md(graph_context=graph)
            self.assertIn("Graph RAG output is unavailable or empty", text)
            self.assertNotIn("shared concepts", text)


class TestMissingData(unittest.TestCase):

    def test_4_missing_bibliographic_metadata(self):
        text = md()
        # Crossref record: no year, no authors.
        self.assertIn("Not available", text)
        refs = build_references(make_papers())["entries"]
        crossref = refs[2]
        reference = ra._reference_text(crossref)
        self.assertIn("Not available", reference)
        self.assertIn("Abstract available: no", reference)

    def test_7_empty_paper_list(self):
        text = md(papers=[])
        self.assertIn("No papers were retrieved.", text)
        self.assertIn("Retrieved papers", text)

    def test_papers_none_and_junk_entries(self):
        self.assertIn("No papers were retrieved.", md(papers=None))
        text = md(papers=[None, "x", 5, make_papers()[0]])
        self.assertIn("Federated Averaging", text)

    def test_10_optional_analysis_fields_absent(self):
        for analysis in ({}, {"summary": "Only a summary."},
                         {"status": "success", "evidence_linked_claims":
                          None}):
            text = md(analysis=analysis, plan={}, graph_context=None)
            self.assertIn("10. References and Source Index", text)
        with tempfile.TemporaryDirectory() as tmp:
            result = generate_research_report(
                QUESTION, None, [], None, {}, output_dir=tmp)
            self.assertEqual(result["errors"], [])
            self.assertTrue(Path(result["docx_path"]).exists())

    def test_insufficient_evidence_analysis(self):
        analysis = {"status": "insufficient_evidence",
                    "summary": "No abstracts.",
                    "evidence_linked_claims": []}
        text = md(analysis=analysis)
        self.assertIn("insufficient evidence", text.lower())
        self.assertIn("no findings are concluded", text)

    def test_no_contradiction_is_stated(self):
        text = md()
        self.assertIn("did not identify a contradiction", text)

    def test_contradiction_is_reported(self):
        analysis = make_analysis()
        analysis["divergent_findings_and_contradictions"] = [{
            "topic": "Noise level", "finding_a": "Low noise works.",
            "paper_a_id": PAPER_A, "finding_b": "High noise needed.",
            "paper_b_id": PAPER_B, "explanation": "They disagree."}]
        text = md(analysis=analysis)
        self.assertIn("Noise level", text)
        self.assertIn("Low noise works.", text)
        self.assertNotIn("did not identify a contradiction", text)

    def test_reported_vs_inferred_gaps_are_distinguished(self):
        text = md()
        self.assertIn("Reported in cited abstract(s)", text)
        self.assertIn("Inferred by the model", text)


class TestVerificationHandling(unittest.TestCase):

    def test_5_all_statuses_are_preserved(self):
        claims = [
            make_claim("CLM-001", "SUPPORTED", excerpt="keeps data"),
            make_claim("CLM-002", "PARTIALLY_SUPPORTED"),
            make_claim("CLM-003", "UNSUPPORTED"),
            make_claim("CLM-004", "UNVERIFIABLE", confidence=0.0),
            make_claim("CLM-005", "SOME_NEW_STATUS"),
        ]
        text = md(analysis=make_analysis(claims))
        for status in ("SUPPORTED", "PARTIALLY_SUPPORTED",
                       "UNSUPPORTED", "UNVERIFIABLE",
                       "SOME_NEW_STATUS"):
            self.assertIn(status, text)
        rows = ra._claim_rows(make_analysis(claims), build_references(
            make_papers()))
        self.assertEqual([r["status"] for r in rows], [
            "SUPPORTED", "PARTIALLY_SUPPORTED", "UNSUPPORTED",
            "UNVERIFIABLE", "SOME_NEW_STATUS"])
        # Everything except SUPPORTED needs manual review.
        self.assertEqual([r["needs_review"] for r in rows],
                         [False, True, True, True, True])

    def test_unverifiable_is_never_shown_as_supported(self):
        analysis = make_analysis([make_claim("CLM-001", "UNVERIFIABLE")])
        row = ra._claim_rows(analysis, build_references(make_papers()))[0]
        self.assertEqual(row["status"], "UNVERIFIABLE")
        self.assertIn("MANUAL REVIEW", row["notes"])

    def test_missing_verification_is_not_verified(self):
        analysis = make_analysis([make_claim("CLM-001", None)],
                                 verified=False)
        text = md(analysis=analysis)
        self.assertIn("NOT VERIFIED", text)
        self.assertIn("Verification did NOT run", text)
        row = ra._claim_rows(analysis, build_references(make_papers()))[0]
        self.assertTrue(row["needs_review"])
        self.assertEqual(row["confidence"], "—")

    def test_excerpt_match_is_not_semantic_verification(self):
        # The Analysis Agent matched the excerpt, but no verifier result:
        analysis = make_analysis([make_claim("CLM-001", None)],
                                 verified=False)
        text = md(analysis=analysis)
        self.assertIn("does not show that the abstract supports the "
                      "whole claim", text)
        row = ra._claim_rows(analysis, build_references(make_papers()))[0]
        self.assertEqual(row["status"], "NOT VERIFIED")
        self.assertIn("| NOT VERIFIED |", text)
        self.assertIn("| SUPPORTED | 0 |", text)

    def test_confidence_is_not_called_a_probability(self):
        text = md()
        self.assertIn("not a calibrated probability", text)

    def test_6_quota_error_is_reported(self):
        analysis = make_analysis([
            make_claim("CLM-001", "UNVERIFIABLE", confidence=0.0)])
        analysis["verification_summary"] = {
            "service_status": "partial",
            "service_note": "Gemini daily quota exhausted for model X.",
            "total_claims": 1}
        text = md(analysis=analysis)
        self.assertIn("PARTIAL", text)
        self.assertIn("daily quota exhausted", text)

    def test_6_verification_failure_passed_from_app(self):
        analysis = make_analysis([make_claim("CLM-001", None)],
                                 verified=False)
        text = md(analysis=analysis,
                  verification_error="RuntimeError: 429 quota")
        self.assertIn("Verification service status: FAILED", text)
        self.assertIn("429 quota", text)
        self.assertIn("NOT VERIFIED", text)

    def test_scope_limitation_is_preserved(self):
        text = md()
        self.assertIn("Only the structured evidence-linked claims", text)

    def test_status_counts_in_summary_section(self):
        claims = [make_claim("CLM-001", "SUPPORTED", excerpt="keeps data"),
                  make_claim("CLM-002", "UNSUPPORTED"),
                  make_claim("CLM-003", "UNSUPPORTED")]
        blocks = build_report_blocks(
            QUESTION, make_plan(), make_papers(), [],
            make_analysis(claims), generated_at=TODAY)
        table = [b for b in blocks if b["type"] == "table"
                 and b["headers"] == ["Measure", "Count"]][0]
        counts = {r[0]: r[1] for r in table["rows"]}
        self.assertEqual(counts["Total claims reviewed"], "3")
        self.assertEqual(counts["SUPPORTED"], "1")
        self.assertEqual(counts["UNSUPPORTED"], "2")
        self.assertEqual(counts["Flagged for manual review"], "2")

    def test_input_analysis_is_not_modified(self):
        analysis = make_analysis()
        snapshot = copy.deepcopy(analysis)
        md(analysis=analysis)
        self.assertEqual(analysis, snapshot)


class TestReferences(unittest.TestCase):

    def test_8_duplicates_are_merged_by_doi_and_title(self):
        papers = make_papers() + [
            # same DOI as paper A but a different OpenAlex-style ID
            {"id": "https://openalex.org/W77", "title": "Federated "
             "Averaging (duplicate)", "doi": "https://doi.org/10.1234/A",
             "url": "https://doi.org/10.1234/A", "abstract": "x"},
            # same title as paper B, different ID, no DOI
            {"id": "https://example.org/other", "title": "SECURE "
             "aggregation!", "abstract": "y"},
        ]
        refs = build_references(papers)
        self.assertEqual(len(refs["entries"]), 3)
        # Aliases resolve to the same reference number.
        self.assertEqual(refs["lookup"]["https://openalex.org/W77"], 1)
        self.assertEqual(refs["lookup"]["https://example.org/other"], 2)
        text = md(papers=papers)
        self.assertEqual(text.count("Paper ID:"), 3)

    def test_first_record_keeps_its_data_and_gaps_are_filled(self):
        papers = [
            {"id": "a", "title": "Same Title", "year": None,
             "authors": [], "abstract": ""},
            {"id": "b", "title": "Same Title", "year": 2020,
             "authors": ["Zed"], "abstract": "Real abstract."},
        ]
        entry = build_references(papers)["entries"][0]
        self.assertEqual(entry["year"], 2020)
        self.assertEqual(entry["authors"], ["Zed"])
        self.assertEqual(entry["abstract"], "Real abstract.")

    def test_claim_citations_resolve_to_reference_numbers(self):
        claim = make_claim("CLM-001", "SUPPORTED", excerpt="keeps data")
        claim["paper_ids"] = [PAPER_B]
        claim["invalid_paper_ids"] = ["https://openalex.org/W404"]
        row = ra._claim_rows(
            make_analysis([claim]), build_references(make_papers()))[0]
        self.assertIn("[2]", row["papers"])
        self.assertIn("unresolved ID: https://openalex.org/W404",
                      row["papers"])

    def test_abstract_not_available_text_is_not_an_abstract(self):
        papers = [{"id": "a", "title": "T",
                   "abstract": "Abstract not available"}]
        self.assertEqual(
            build_references(papers)["entries"][0]["abstract"], "")


class TestSpecialCharacters(unittest.TestCase):

    NASTY = ("Pipes | in text,\nnewline here\r\n* star _under_ "
             "`tick` <b>tag</b> [link](x) # hash \x00 \x0b end")

    def test_9_markdown_table_stays_valid(self):
        claim = make_claim("CLM-001", "SUPPORTED", excerpt="x")
        claim["claim"] = self.NASTY
        claim["evidence_excerpt"] = self.NASTY
        claim["verification"]["evidence_excerpt"] = self.NASTY
        text = md(analysis=make_analysis([claim]))
        # Every table row has the same number of unescaped pipes as its
        # header: no stray "|" or newline broke the table.
        table_lines = [l for l in text.splitlines()
                       if l.startswith("| ") and "Claim ID" in l]
        header = table_lines[0]
        start = text.index(header)
        rows = []
        for line in text[start:].splitlines():
            if not line.startswith("|"):
                break
            rows.append(line)
        import re
        pipe = lambda s: len(re.findall(r"(?<!\\)\|", s))
        self.assertEqual({pipe(r) for r in rows}, {pipe(header)})
        self.assertNotIn("\x00", text)
        self.assertNotIn("\x0b", text)
        self.assertIn("Pipes \\| in text", text)
        self.assertNotIn("<b>", text)

    def test_9_docx_survives_special_characters(self):
        claim = make_claim("CLM-001", "SUPPORTED", excerpt="x")
        claim["claim"] = self.NASTY
        claim["verification"]["evidence_excerpt"] = self.NASTY
        papers = make_papers()
        papers[0]["title"] = "Title with \x00 control & <xml> \"quotes\""
        with tempfile.TemporaryDirectory() as tmp:
            result = generate_research_report(
                QUESTION, make_plan(), papers, make_graph(),
                make_analysis([claim]), output_dir=tmp)
            self.assertEqual(result["errors"], [])
            doc = Document(result["docx_path"])
            cells = " ".join(c.text for t in doc.tables for r in t.rows
                             for c in r.cells)
            self.assertIn("Pipes | in text, newline here", cells)

    def test_multiline_excerpt_is_single_line(self):
        self.assertEqual(ra._clean("a\n\nb\r\nc\t d"), "a b c d")

    def test_urls_become_valid_markdown_links(self):
        out = ra._md_inline("See https://example.org/a_(b)_c now.")
        self.assertIn("(https://example.org/a_%28b%29_c)", out)
        # Trailing punctuation and an unbalanced ")" are not in the URL.
        out = ra._md_inline("(see https://example.org/x).")
        self.assertIn("(https://example.org/x)", out)
        self.assertNotIn("x%29", out)

    def test_unicode_is_preserved(self):
        papers = [{"id": "u", "title": "Étude de l'apprentissage — 联邦学习",
                   "abstract": "résumé"}]
        text = md(papers=papers)
        self.assertIn("联邦学习", text)


class TestAnalysisCompatibility(unittest.TestCase):

    def test_works_with_real_agent_outputs(self):
        """Report accepts real analysis + verification agent output."""
        from agents.analysis_agent import (
            build_corpus_index, validate_and_enrich_analysis)
        from agents.verification_agent import verify_analysis
        from unittest.mock import MagicMock
        import json

        papers = make_papers()
        raw = {
            "summary": "Summary.",
            "key_themes_and_consensus": [],
            "divergent_findings_and_contradictions": [],
            "methodologies_and_benchmarks": [],
            "limitations_and_research_gaps": [],
            "evidence_linked_claims": [{
                "claim_id": "CLM-001",
                "claim": "FL keeps data on device.",
                "support_type": "source_supported",
                "paper_ids": [PAPER_A],
                "evidence_excerpt": "keeps data on device"}],
        }
        analysis = validate_and_enrich_analysis(
            raw, build_corpus_index(papers))

        client = MagicMock()
        reply = MagicMock()
        reply.text = json.dumps({
            "status": "SUPPORTED", "confidence": 0.9,
            "justification": "Stated directly.",
            "evidence": [{"paper_id": PAPER_A,
                          "excerpt": "keeps data on device"}],
            "warning": ""})
        client.models.generate_content.return_value = reply
        verified = verify_analysis(analysis, papers, client=client)

        text = md(analysis=verified, papers=papers)
        self.assertIn("SUPPORTED", text)
        self.assertIn("keeps data on device", text)
        self.assertIn("Stated directly.", text)
        self.assertIn("[1]", text)
        self.assertIn("| SUPPORTED | 0.90 |", text)


if __name__ == "__main__":
    unittest.main()